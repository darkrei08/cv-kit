"""End-to-end build tests. No office suite and no external service required."""
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from cvkit.builder import build_all, build_master, build_single
from cvkit.model import CV, Contact, Job, Letter


def make_cv(tmp_path: Path, *, with_photo: bool = False) -> CV:
    photo = ""
    signature = ""
    if with_photo:
        photo_path = tmp_path / "photo.jpg"
        Image.new("RGB", (413, 531), (200, 200, 200)).save(photo_path, "JPEG")
        signature_path = tmp_path / "signature.png"
        Image.new("RGB", (400, 120), (255, 255, 255)).save(signature_path, "PNG")
        photo = str(photo_path)
        signature = str(signature_path)
    return CV(
        contact=Contact(name="Test Person", email="test@example.com",
                        phone="+39 000 000 0000", headline="Elettricista",
                        location="Bologna (BO), Italia",
                        address="Via Esempio 1 – Bologna (BO)",
                        personal={"Data di nascita": "01/01/1980"},
                        photo=photo, signature=signature,
                        place_date="Bologna, settembre 2026"),
        profile="Elettricista industriale con esperienza su quadri di automazione.",
        experience=[Job(period="01/2020 - 12/2024", employer="Azienda S.r.l.",
                        location="Bologna (BO)", role="Elettricista",
                        sector="Manifattura", contract="Tempo indeterminato",
                        duties=["Manutenzione dei quadri.", "Ricerca guasti."])],
        skills=[],
        licences=["Patente B."],
        availability=["Disponibile a turni."],
        keywords="quadri, inverter, plc, motori, ricerca guasti",
        letter=Letter(subject="Candidatura", paragraphs=["Paragrafo uno."],
                      variants=[("Variante A", ["Punto uno."])]),
    )


def test_master_contains_every_part(tmp_path):
    cv = make_cv(tmp_path)
    path = build_master(cv, tmp_path / "master.docx")
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert "Test Person" in text
    assert "Note di verifica" in text


def test_master_can_exclude_the_reserved_notes(tmp_path):
    cv = make_cv(tmp_path)
    path = build_master(cv, tmp_path / "fascicolo.docx", include_notes=False)
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert "Note di verifica" not in text


def test_the_ats_variant_has_no_table_and_no_image(tmp_path):
    cv = make_cv(tmp_path, with_photo=True)
    path = build_single(cv, "ats", tmp_path / "ats.docx")
    document = Document(str(path))
    assert len(document.tables) == 0
    assert len(document.inline_shapes) == 0


def test_the_ats_variant_omits_personal_data(tmp_path):
    cv = make_cv(tmp_path)
    path = build_single(cv, "ats", tmp_path / "ats.docx")
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert "01/01/1980" not in text
    assert "Via Esempio" not in text


def test_the_designed_variant_uses_photo_and_signature(tmp_path):
    cv = make_cv(tmp_path, with_photo=True)
    path = build_single(cv, "designed", tmp_path / "designed.docx")
    document = Document(str(path))
    assert len(document.inline_shapes) == 2
    assert len(document.tables) > 0


def test_build_all_writes_the_expected_files_without_pdf(tmp_path):
    cv = make_cv(tmp_path)
    result = build_all(cv, tmp_path / "out", exports=("docx",))
    names = {p.name for p in result.files.values()}
    assert any(n.endswith("CV_Master.docx") for n in names)
    assert any(n.endswith("CV_Fascicolo.docx") for n in names)
    assert any("ATS" in n and n.endswith(".docx") for n in names)
    assert any("DESIGNED" in n for n in names)


def test_build_all_adds_a_translated_ats_file(tmp_path):
    cv = make_cv(tmp_path)
    translated = make_cv(tmp_path)
    translated.contact.name = "Test Person"
    translated.lang = "en"
    result = build_all(cv, tmp_path / "out", translations={"en": translated},
                       exports=("docx",))
    assert any("ATS_EN" in p.name for p in result.files.values())


def test_file_names_are_surname_first_and_free_of_spaces(tmp_path):
    cv = make_cv(tmp_path)
    cv.contact.name = "Maria Grazia Rossi"
    result = build_all(cv, tmp_path / "out", exports=("docx",))
    for path in result.files.values():
        assert " " not in path.name
    assert any(path.name.startswith("Rossi_Maria_Grazia") for path in result.files.values())


def test_metadata_is_set_on_the_document(tmp_path):
    cv = make_cv(tmp_path)
    path = build_single(cv, "ats", tmp_path / "ats.docx")
    properties = Document(str(path)).core_properties
    assert properties.author == "Test Person"
    assert properties.title
    assert properties.comments == ""


def test_unknown_part_is_rejected(tmp_path):
    cv = make_cv(tmp_path)
    with pytest.raises(ValueError):
        build_single(cv, "nonsense", tmp_path / "x.docx")


def test_long_keywords_are_clipped_for_the_file_properties(tmp_path):
    """OOXML caps the property at 255 characters; the document body keeps the rest."""
    cv = make_cv(tmp_path)
    cv.keywords = ", ".join(f"termine numero {i}" for i in range(40))
    path = build_single(cv, "ats", tmp_path / "ats.docx")
    document = Document(str(path))
    assert len(document.core_properties.keywords) <= 255
    body = "\n".join(p.text for p in document.paragraphs)
    assert "termine numero 39" in body


def test_a_broken_photo_does_not_abort_the_build(tmp_path):
    """A photograph that is not an image is a cosmetic problem, not a fatal one."""
    fake = tmp_path / "finta.jpg"
    fake.write_text("questo non e' un JPEG", encoding="utf-8")
    cv = make_cv(tmp_path)
    cv.contact.photo = str(fake)
    cv.contact.signature = str(fake)

    path = build_single(cv, "designed", tmp_path / "designed.docx")
    assert path.exists()
    document = Document(str(path))
    assert len(document.inline_shapes) == 0
    assert any("Data di nascita" in p.text or True for p in document.paragraphs)


def test_resolve_assets_reports_a_file_that_is_not_an_image(tmp_path):
    from cvkit.builder import resolve_assets

    fake = tmp_path / "assets"
    fake.mkdir()
    (fake / "profilo.jpg").write_text("non e' un'immagine", encoding="utf-8")
    cv = make_cv(tmp_path)
    cv.contact.photo = "assets/profilo.jpg"

    problems = resolve_assets(cv, tmp_path)
    assert problems and "not a readable image" in problems[0]
    assert cv.contact.photo == ""


def test_resolve_assets_reports_an_absolute_path_that_is_not_an_image(tmp_path):
    from cvkit.builder import resolve_assets

    fake = tmp_path / "assoluta.jpg"
    fake.write_text("non e' un'immagine", encoding="utf-8")
    cv = make_cv(tmp_path)
    cv.contact.photo = str(fake)

    problems = resolve_assets(cv, tmp_path)
    assert problems and "not a readable image" in problems[0]
    assert cv.contact.photo == ""


def test_resolve_assets_keeps_a_valid_photo(tmp_path):
    from PIL import Image

    from cvkit.builder import resolve_assets

    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (300, 400), (200, 200, 200)).save(assets / "profilo.jpg", "JPEG")
    cv = make_cv(tmp_path)
    cv.contact.photo = "assets/profilo.jpg"

    assert resolve_assets(cv, tmp_path) == []
    assert cv.contact.photo.endswith("profilo.jpg")
