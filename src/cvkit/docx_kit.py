# -*- coding: utf-8 -*-
"""Low-level Word primitives.

Everything here is a thin, explicit wrapper over python-docx and raw OOXML. It is
kept separate from the CV builders on purpose: these functions are the ones that
are hard to get right (hanging indents, right-aligned tabs, paragraph borders, cell
shading, fixed table layout), and they are worth testing on their own.
"""
from __future__ import annotations

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageFilter

from .theme import Palette


def el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for key, value in attrs.items():
        e.set(qn(key), str(value))
    return e


# ---------------------------------------------------------------------------
# Page
# ---------------------------------------------------------------------------
def setup_page(section, page) -> None:
    section.page_width = Mm(page.width_mm)
    section.page_height = Mm(page.height_mm)
    section.top_margin = Mm(page.top)
    section.right_margin = Mm(page.right)
    section.bottom_margin = Mm(page.bottom)
    section.left_margin = Mm(page.left)
    section.header_distance = Mm(9)
    section.footer_distance = Mm(9)


def set_doc_defaults(doc, palette: Palette, size: float, lang: str = "it-IT") -> None:
    style = doc.styles["Normal"]
    style.font.name = palette.body_font
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(palette.ink)
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.06
    _force_font(style, palette.body_font)
    rpr = style.element.get_or_add_rPr()
    rpr.append(el("w:lang", **{"w:val": lang, "w:eastAsia": lang, "w:bidi": lang}))


def _force_font(style, font_name: str) -> None:
    """Set the font on every script slot, or Word substitutes for CJK/complex text."""
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = el("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font_name)


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------
def make_style(doc, name: str, size: float, *, bold: bool = False, italic: bool = False,
               color: str = "000000", font: str = "Arial", before: float = 0,
               after: float = 2, line: float = 1.06, allcaps: bool = False,
               indent: float | None = None, hanging: float | None = None,
               keep_next: bool = False, base: str = "Normal"):
    if name in [s.name for s in doc.styles]:
        return doc.styles[name]
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles[base]
    style.quick_style = True
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.all_caps = allcaps
    style.font.name = font
    style.font.color.rgb = RGBColor.from_string(color)
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if indent is not None:
        pf.left_indent = Mm(indent)
    if hanging is not None:
        pf.first_line_indent = Mm(-hanging)
    _force_font(style, font)
    return style


# ---------------------------------------------------------------------------
# Paragraphs
# ---------------------------------------------------------------------------
def para(parent, text: str = "", style: str | None = None, *, size: float | None = None,
         bold: bool | None = None, italic: bool | None = None, color: str | None = None,
         align=None, before: float | None = None, after: float | None = None,
         line: float | None = None, font: str | None = None,
         keep_next: bool | None = None, indent: float | None = None,
         hanging: float | None = None):
    p = parent.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        if size is not None:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if italic is not None:
            run.font.italic = italic
        if color is not None:
            run.font.color.rgb = RGBColor.from_string(color)
        if font is not None:
            run.font.name = font
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    for attr, value in (("space_before", before), ("space_after", after),
                        ("line_spacing", line)):
        if value is not None:
            setattr(pf, attr, Pt(value) if attr != "line_spacing" else value)
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if indent is not None:
        pf.left_indent = Mm(indent)
    if hanging is not None:
        pf.first_line_indent = Mm(-hanging)
    return p


def rich(parent, chunks, style: str | None = None, *, align=None, before: float | None = None,
         after: float | None = None, line: float | None = None,
         indent: float | None = None, hanging: float | None = None,
         keep_next: bool | None = None):
    """One paragraph built from several runs.

    ``chunks`` is a sequence of ``(text, {bold, italic, color, size, font})``.
    """
    p = parent.add_paragraph(style=style)
    for text, fmt in chunks:
        run = p.add_run(text)
        run.font.bold = fmt.get("bold", False)
        run.font.italic = fmt.get("italic", False)
        if "size" in fmt:
            run.font.size = Pt(fmt["size"])
        if "color" in fmt:
            run.font.color.rgb = RGBColor.from_string(fmt["color"])
        if "font" in fmt:
            run.font.name = fmt["font"]
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    for attr, value in (("space_before", before), ("space_after", after),
                        ("line_spacing", line)):
        if value is not None:
            setattr(pf, attr, Pt(value) if attr != "line_spacing" else value)
    if indent is not None:
        pf.left_indent = Mm(indent)
    if hanging is not None:
        pf.first_line_indent = Mm(-hanging)
    if keep_next is not None:
        pf.keep_with_next = keep_next
    return p


def bullet(parent, text: str, *, style: str = "bullet", marker: str = "\u2022  ",
           size: float | None = None, lead: str | None = None, lead_color: str = "404040",
           marker_color: str = "404040"):
    """A bullet with a literal glyph, a hanging indent and a bold lead-in.

    A literal marker rather than Word numbering: it survives every converter, every
    Word version and every plain-text extraction, which is exactly what a CV needs.
    """
    p = parent.add_paragraph(style=style)
    glyph = p.add_run(marker)
    glyph.font.bold = True
    glyph.font.color.rgb = RGBColor.from_string(marker_color)
    if lead:
        head = p.add_run(lead)
        head.font.bold = True
        head.font.color.rgb = RGBColor.from_string(lead_color)
        p.add_run(" \u2013 " + text)
    else:
        p.add_run(text)
    if size is not None:
        for run in p.runs:
            run.font.size = Pt(size)
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def tab_right(p, position_mm: float):
    p.paragraph_format.tab_stops.add_tab_stop(Mm(position_mm), WD_TAB_ALIGNMENT.RIGHT)
    return p


def tab_left(p, position_mm: float):
    p.paragraph_format.tab_stops.add_tab_stop(Mm(position_mm), WD_TAB_ALIGNMENT.LEFT)
    return p


# ---------------------------------------------------------------------------
# Borders, shading, cells
# ---------------------------------------------------------------------------
def para_border(p, edge: str = "bottom", color: str = "8C8C8C", size: int = 6,
                space: int = 3):
    ppr = p._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    borders.append(el("w:" + edge, **{"w:val": "single", "w:sz": size,
                                      "w:space": space, "w:color": color}))
    return p


def para_shade(p, fill: str):
    p._p.get_or_add_pPr().append(
        el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))
    return p


def cell_shade(cell, fill: str):
    cell._tc.get_or_add_tcPr().append(
        el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))
    return cell


def clear_cell(cell):
    """Remove the empty paragraph python-docx creates inside a new cell."""
    for p in list(cell.paragraphs):
        p._p.getparent().remove(p._p)
    return cell


def ensure_cell_content(table):
    """Guarantee every cell holds at least one block element.

    A ``w:tc`` with no paragraph or table is invalid OOXML. Word does not repair it:
    it refuses to open the document and reports that the file may be corrupted, which
    points nowhere near the real cause. Every table that had a cell emptied by
    :func:`clear_cell` must pass through here.
    """
    for row in table.rows:
        for cell in row.cells:
            if not cell.paragraphs and not cell.tables:
                cell.add_paragraph()
    return table


def table_borders(table, color: str = "C7D3DE", size: int = 4,
                  edges=("top", "left", "bottom", "right", "insideH", "insideV")):
    borders = el("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in edges:
            borders.append(el("w:" + edge, **{"w:val": "single", "w:sz": size,
                                              "w:space": 0, "w:color": color}))
        else:
            borders.append(el("w:" + edge, **{"w:val": "none", "w:sz": 0,
                                              "w:space": 0, "w:color": "auto"}))
    table._tbl.tblPr.append(borders)
    return table


def fixed_layout(table):
    """Stop Word from resizing columns to fit the content."""
    table._tbl.tblPr.append(el("w:tblLayout", **{"w:type": "fixed"}))
    return table


def keep_rows_together(table):
    for row in table.rows:
        row._tr.get_or_add_trPr().append(el("w:cantSplit"))
    return table


def table_widths(table, widths_mm) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths_mm):
            cell.width = Mm(width)


def header_row(cells, widths_mm, labels, fill: str, size: float = 8.5) -> None:
    """Shade a table's first row and write white bold labels into it."""
    for cell, width, label in zip(cells, widths_mm, labels):
        cell.width = Mm(width)
        cell_shade(cell, fill)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        run.font.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string("FFFFFF")


def body_row(cells, widths_mm, values, size: float = 8.5, fill: str | None = None) -> None:
    for cell, width, value in zip(cells, widths_mm, values):
        cell.width = Mm(width)
        if fill:
            cell_shade(cell, fill)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.space_before = Pt(1)
        paragraph.add_run(value).font.size = Pt(size)


# ---------------------------------------------------------------------------
# Images
# ---------------------------------------------------------------------------
def prepare_photo(src: str, dst: str, size: tuple[int, int] = (413, 531),
                  sharpen: bool = True) -> str:
    """Crop to a portrait ratio, upscale with Lanczos, sharpen, save at 300 dpi.

    A small source photo cannot be invented into a large one, but at 30-35 mm wide
    an upscaled 413x531 file prints at 300 dpi and looks correct.
    """
    image = Image.open(src).convert("RGB")
    width, height = image.size
    target_ratio = size[0] / size[1]
    if width / height > target_ratio:
        new_width = int(height * target_ratio)
        left = (width - new_width) // 2
        image = image.crop((left, 0, left + new_width, height))
    else:
        new_height = int(width / target_ratio)
        top = (height - new_height) // 2
        image = image.crop((0, top, width, top + new_height))
    image = image.resize(size, Image.LANCZOS)
    if sharpen:
        image = image.filter(ImageFilter.UnsharpMask(radius=2, percent=110, threshold=3))
    image.save(dst, "JPEG", quality=94, dpi=(300, 300))
    return dst


def prepare_signature(src: str, dst: str) -> str:
    """Flatten a scanned signature onto white so it prints cleanly."""
    image = Image.open(src).convert("L")
    image = image.point(lambda v: 255 if v > 170 else int(v * 0.35)).convert("RGB")
    image.save(dst, "PNG", dpi=(300, 300))
    return dst


def centred_image(doc, path: str, width_mm: float):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Mm(width_mm))
    return p
