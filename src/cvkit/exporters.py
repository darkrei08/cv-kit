"""Export a built document to the formats a job application actually needs.

PDF export is the one genuinely platform-specific step in the whole pipeline,
because there is no pure-Python renderer that reproduces Word pagination. Three
backends are supported and probed in order of fidelity:

``word``
    Microsoft Word through its COM interface, driven from PowerShell on Windows.
    Highest fidelity: the PDF pagination is identical to what Word prints.

``docx2pdf``
    The ``docx2pdf`` package (Word on Windows, Word on macOS). Same engine, one
    less moving part, but an extra dependency.

``libreoffice``
    ``soffice --headless --convert-to pdf``. Available on every platform and the
    only option on Linux. Pagination is very close but not guaranteed identical to
    Word, which is why :mod:`cvkit.qa` re-measures the result.

Text export uses ``pandoc`` when present and falls back to a small pure-Python
extraction, so a machine without pandoc still produces the file that web forms ask
for.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path

# ---------------------------------------------------------------------------
# Locating the tools
# ---------------------------------------------------------------------------
WINDOWS = os.name == "nt"


def _which(*candidates: str) -> str | None:
    """Find an executable. Only files count: a directory with the same name is not
    a program, and returning one produces a confusing failure much later."""
    for candidate in candidates:
        found = shutil.which(candidate)
        if found and os.path.isfile(found):
            return found
        if WINDOWS:
            for root in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                         os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
                         os.path.expanduser("~\\AppData\\Local")):
                guess = Path(root) / candidate
                if guess.is_file():
                    return str(guess)
    return None


def find_libreoffice() -> str | None:
    return _which(
        "soffice", "soffice.exe", "libreoffice",
        r"LibreOffice\program\soffice.exe",
        r"Programs\LibreOffice\program\soffice.exe",
    )


def find_pandoc() -> str | None:
    return _which("pandoc", "pandoc.exe", r"Pandoc\pandoc.exe")


def find_powershell() -> str | None:
    if not WINDOWS:
        return None
    return _which("powershell", "powershell.exe", "pwsh", "pwsh.exe")


def word_available() -> bool:
    if not WINDOWS:
        return False
    for root in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                 os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
        for candidate in ("Microsoft Office/root/Office16/WINWORD.EXE",
                          "Microsoft Office/root/Office15/WINWORD.EXE"):
            if (Path(root) / candidate).exists():
                return True
    return False


def docx2pdf_available() -> bool:
    try:
        import docx2pdf  # noqa: F401
    except Exception:
        return False
    return WINDOWS or os.uname().sysname == "Darwin" if hasattr(os, "uname") else WINDOWS


def backends() -> dict[str, bool]:
    """Report which export backends this machine can use."""
    return {
        "word": word_available() and find_powershell() is not None,
        "docx2pdf": docx2pdf_available(),
        "libreoffice": find_libreoffice() is not None,
        "pandoc": find_pandoc() is not None,
    }


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
POWERSHELL_SCRIPT = r"""param([Parameter(Mandatory = $true)][string]$In,
                             [Parameter(Mandatory = $true)][string]$Out)
$ErrorActionPreference = "Stop"
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {
    $doc = $word.Documents.Open($In, $false, $true)
    $doc.SaveAs([ref]$Out, [ref]17)
    $doc.Close([ref]0)
} finally {
    $word.Quit()
    [System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null
}
"""


def _export_with_word(docx: Path, pdf: Path) -> None:
    powershell = find_powershell()
    if not powershell:
        raise RuntimeError("PowerShell is not available")
    # Absolute paths: Word resolves relative paths against its own working
    # directory, not ours, and fails with a confusing "file not found".
    docx = docx.resolve()
    pdf = pdf.resolve()
    with tempfile.TemporaryDirectory() as tmp:
        script = Path(tmp) / "docx_to_pdf.ps1"
        script.write_text(POWERSHELL_SCRIPT, encoding="utf-8")
        result = subprocess.run(
            [powershell, "-NoProfile", "-ExecutionPolicy", "Bypass",
             "-File", str(script), "-In", str(docx), "-Out", str(pdf)],
            capture_output=True, text=True, timeout=600,
        )
    if result.returncode != 0 or not pdf.exists():
        raise RuntimeError(f"Word export failed: {result.stderr.strip()[:400]}")


def _export_with_docx2pdf(docx: Path, pdf: Path) -> None:
    from docx2pdf import convert  # type: ignore
    convert(str(docx), str(pdf))
    if not pdf.exists():
        raise RuntimeError("docx2pdf produced no file")


def _export_with_libreoffice(docx: Path, pdf: Path) -> None:
    soffice = find_libreoffice()
    if not soffice:
        raise RuntimeError("LibreOffice is not installed")
    docx = docx.resolve()
    pdf = pdf.resolve()
    pdf.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=str(pdf.parent)) as tmp:
        result = subprocess.run(
            [soffice, "--headless", "--norestore", "--convert-to", "pdf",
             "--outdir", tmp, str(docx)],
            capture_output=True, text=True, timeout=600,
        )
        produced = Path(tmp) / (docx.stem + ".pdf")
        if not produced.exists():
            raise RuntimeError(
                "LibreOffice produced no PDF: "
                f"{result.stdout.strip()} {result.stderr.strip()}"[:400])
        shutil.move(str(produced), str(pdf))


_BACKEND_ORDER = ("word", "docx2pdf", "libreoffice")
_BACKEND_FUNCS = {
    "word": _export_with_word,
    "docx2pdf": _export_with_docx2pdf,
    "libreoffice": _export_with_libreoffice,
}


def export_pdf(docx: str | Path, pdf: str | Path | None = None,
               backend: str = "auto") -> Path:
    """Convert a ``.docx`` to PDF, choosing the best available engine.

    Pass ``backend`` to force one (``"word"``, ``"docx2pdf"``, ``"libreoffice"``).
    Paths are resolved to absolute before any external tool sees them: a relative
    path is interpreted by the tool's own working directory, which is a frequent
    source of "file not found" errors that look like a missing document.
    """
    docx = Path(docx).resolve()
    pdf = Path(pdf).resolve() if pdf else docx.with_suffix(".pdf")
    if not docx.exists():
        raise FileNotFoundError(docx)

    if backend != "auto":
        if backend not in _BACKEND_FUNCS:
            raise ValueError(f"unknown backend '{backend}'; "
                             f"choose one of {', '.join(_BACKEND_ORDER)}")
        _BACKEND_FUNCS[backend](docx, pdf)
        return pdf

    problems = []
    for name in _BACKEND_ORDER:
        if not backends().get(name):
            continue
        try:
            _BACKEND_FUNCS[name](docx, pdf)
            return pdf
        except Exception as exc:  # try the next engine
            problems.append(f"{name}: {exc}")
    raise RuntimeError(
        "no PDF backend worked.\n"
        + ("\n".join(f"  - {p}" for p in problems) if problems else "")
        + "\nInstall Microsoft Word (Windows), or LibreOffice:\n"
          "  Windows  winget install TheDocumentFoundation.LibreOffice\n"
          "  macOS    brew install --cask libreoffice\n"
          "  Debian   sudo apt install libreoffice-writer\n"
          "  Fedora   sudo dnf install libreoffice-writer\n"
          "  Alpine   apk add libreoffice-writer\n"
          "  Arch     sudo pacman -S libreoffice-fresh"
    )


# ---------------------------------------------------------------------------
# Plain text (the field a web form asks you to paste into)
# ---------------------------------------------------------------------------
def export_text(docx: str | Path, txt: str | Path | None = None) -> Path:
    docx = Path(docx)
    txt = Path(txt) if txt else docx.with_suffix(".txt")
    pandoc = find_pandoc()
    if pandoc:
        result = subprocess.run(
            [pandoc, str(docx), "-t", "plain", "--wrap=none", "-o", str(txt)],
            capture_output=True, text=True, timeout=180,
        )
        if result.returncode == 0 and txt.exists():
            return txt
    _text_fallback(docx, txt)
    return txt


def _text_fallback(docx: Path, txt: Path) -> None:
    """Extract paragraphs and table cells without any external tool."""
    from docx import Document

    document = Document(str(docx))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    txt.write_text("\n".join(lines) + "\n", encoding="utf-8")


def export_html_pdf(cv, html_path: str | Path, pdf_path: str | Path | None = None):
    """Render the HTML/CSS variant to PDF with WeasyPrint, when installed."""
    from .html_render import html_to_pdf, render_html

    html_path = Path(html_path)
    render_html(cv, html_path)
    if pdf_path is None:
        return html_path
    return html_to_pdf(html_path, Path(pdf_path))


def doctor() -> str:
    """Human-readable report of what this machine can do."""
    state = backends()
    lines = ["cv-kit environment", ""]
    for name in ("word", "docx2pdf", "libreoffice", "pandoc"):
        lines.append(f"  {'yes' if state[name] else 'no ':>3}  {name}")
    try:
        import docx  # noqa: F401
        lines.append("  yes  python-docx")
    except Exception:
        lines.append("  no   python-docx (required)")
    try:
        import PIL  # noqa: F401
        lines.append("  yes  pillow")
    except Exception:
        lines.append("  no   pillow (required for photo handling)")
    try:
        import pymupdf  # noqa: F401
        lines.append("  yes  pymupdf (quality checks)")
    except Exception:
        lines.append("  no   pymupdf (pip install 'cvkit[qa]')")
    try:
        import weasyprint  # noqa: F401
        lines.append("  yes  weasyprint (HTML variant)")
    except Exception:
        lines.append("  no   weasyprint (pip install 'cvkit[html]')")
    if find_libreoffice():
        lines.append(f"\n  LibreOffice: {find_libreoffice()}")
    if find_pandoc():
        lines.append(f"  pandoc:      {find_pandoc()}")
    return "\n".join(lines)
